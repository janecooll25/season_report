"""Создаёт пример docx-шаблона с плейсхолдерами docxtpl (Jinja2).

Запуск: python scripts/make_sample_template.py
Свой шаблон можно сделать в Word, используя те же {{ ... }} и {% ... %} теги.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = Path("templates/report_template.docx")


def add_table(doc: Document, header: str, list_var: str, first_col: str) -> None:
    doc.add_heading(header, level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = first_col
    hdr[1].text = "Визиты"
    hdr[2].text = "Пользователи"
    hdr[3].text = "Доля, %"

    # docxtpl: тег {%tr ...%} удаляет строку целиком, поэтому for/endfor
    # выносим в отдельные строки-обёртки вокруг строки с данными.
    table.add_row().cells[0].text = "{%%tr for r in %s %%}" % list_var
    data = table.add_row().cells
    data[0].text = "{{ r.name }}"
    data[1].text = "{{ r.visits }}"
    data[2].text = "{{ r.users }}"
    data[3].text = "{{ r.share }}"
    table.add_row().cells[0].text = "{%tr endfor %}"


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    title = doc.add_heading("{{ title }}", level=0)
    title.runs[0].font.size = Pt(22)

    p = doc.add_paragraph()
    p.add_run("Счётчик: ").bold = True
    p.add_run("{{ counter_id }}")
    p.add_run("    Период: ").bold = True
    p.add_run("{{ date_from }} — {{ date_to }}")
    p.add_run("    Сформирован: ").bold = True
    p.add_run("{{ generated_at }}")

    doc.add_heading("Сводка", level=1)
    summary = doc.add_table(rows=6, cols=2)
    summary.style = "Light List Accent 1"
    pairs = [
        ("Визиты", "{{ visits }}"),
        ("Уникальные пользователи", "{{ users }}"),
        ("Просмотры страниц", "{{ pageviews }}"),
        ("Отказы, %", "{{ bounce_rate }}"),
        ("Средняя длительность визита", "{{ avg_duration }}"),
        ("Глубина просмотра", "{{ page_depth }}"),
    ]
    for i, (k, v) in enumerate(pairs):
        summary.rows[i].cells[0].text = k
        summary.rows[i].cells[1].text = v

    add_table(doc, "Источники трафика", "sources", "Источник")
    add_table(doc, "Топ страниц", "top_pages", "Страница")
    add_table(doc, "Устройства", "devices", "Тип устройства")

    doc.add_heading("Выводы", level=1)
    doc.add_paragraph(
        "За период с {{ date_from }} по {{ date_to }} сайт получил {{ visits }} "
        "визитов от {{ users }} пользователей. Доля отказов составила "
        "{{ bounce_rate }}%, средняя глубина просмотра — {{ page_depth }}."
    )

    doc.save(str(OUT))
    print(f"Шаблон создан: {OUT}")


if __name__ == "__main__":
    main()
