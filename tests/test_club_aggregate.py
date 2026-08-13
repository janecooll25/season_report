from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from test_ticket_metrics import _make_raw
from yandex_report.club_aggregate import (
    MAX_CLUBS, aggregate_clubs, build_aggregate_docx, compute_club_metrics,
)
from yandex_report.ticket_metrics import RZ, TicketError


def test_compute_club_metrics_basic():
    m = compute_club_metrics(_make_raw(), capacity=200)
    # 2 матча регулярки + 1 плей-офф = 172+190+265 = 627 проходов
    assert m["att_season"] == 627
    assert m["matches"] == 3
    assert m["income_total"] == 720000  # 200000+220000+300000
    assert m["online_share"] == pytest.approx(300 / 400)  # онлайн 300 из 400


def test_aggregate_has_four_params_with_grades():
    files = [("HC_A.xlsx", _make_raw()), ("HC_B.xlsx", _make_raw())]
    out = aggregate_clubs(files, capacity=200, season_label="2025/2026")
    ws = openpyxl.load_workbook(BytesIO(out)).active
    assert ws.title == "Средние по клубам"
    assert "Клубов в выборке: 2" in ws["A2"].value

    col_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    # 4 параметра-секции
    assert any("Средняя цена билета (регулярный чемпионат)" in str(v) for v in col_a)
    assert any("Доля продаж билетов онлайн" in str(v) for v in col_a)
    assert any("Доля бесплатных билетов" in str(v) for v in col_a)
    assert any("Фактическое отклонение посещаемости" in str(v) for v in col_a)
    # клубы и градация присутствуют
    assert "HC A" in col_a and "HC B" in col_a
    grades = [ws.cell(r, 4).value for r in range(1, ws.max_row + 1)]
    assert any("показатели" in str(g) for g in grades)


def test_aggregate_docx_lists_params_and_places():
    from io import BytesIO as _B

    from docx import Document

    files = [("HC_A.xlsx", _make_raw()), ("HC_B.xlsx", _make_raw())]
    out = build_aggregate_docx(files, capacity=200, season_label="2025/2026")
    doc = Document(_B(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Средняя цена билета" in h for h in headings)
    assert any("Доля продаж билетов онлайн" in h for h in headings)
    # таблицы с колонкой «Место» и клубами, ранжированными по местам
    assert doc.tables
    hdr = [c.text for c in doc.tables[0].rows[0].cells]
    assert hdr == ["Место", "Клуб", "Значение", "Градация"]
    places = [doc.tables[0].rows[i].cells[0].text for i in range(1, 3)]
    assert places == ["1", "2"]


def test_paid_column_as_formula_is_evaluated():
    """Столбец «платные билеты» задан формулой (=всего−беспл.) без кэша значений."""
    wb = openpyxl.load_workbook(BytesIO(_make_raw()))
    rz = wb[RZ]
    for r in (5, 6, 8):  # как в реальных выгрузках: только разовые + беспл., C = K − D
        for c in range(5, 11):  # обнуляем абонементы/бизнес/ложи (E..J)
            rz.cell(r, c).value = 0
        rz.cell(r, 11).value = rz.cell(r, 3).value + rz.cell(r, 4).value  # K = C + D
        rz.cell(r, 3).value = f"=K{r}-D{r}"                               # C — формулой
    buf = BytesIO(); wb.save(buf)

    m = compute_club_metrics(buf.getvalue(), capacity=200)
    # цена регулярки считается (формула C прочитана, а не взята за 0)
    assert m["price_reg"] is not None and m["price_reg"] > 0
    # посещаемость = сумма K: (100+10)+(120+8)+(200+4) = 442
    assert m["att_season"] == 442


def test_to_rub_scales_money_metrics_only():
    """to_rub пересчитывает деньги в рубли, но не трогает доли/проценты."""
    base = compute_club_metrics(_make_raw(), capacity=200)
    conv = compute_club_metrics(_make_raw(), capacity=200, to_rub=27.0)
    assert conv["income_total"] == pytest.approx(base["income_total"] * 27.0)
    assert conv["price_reg"] == pytest.approx(base["price_reg"] * 27.0)
    # доли не денежные — без изменений
    assert conv["free_share"] == base["free_share"]
    assert conv["online_share"] == base["online_share"]


def test_deviation_signed_display_ranked_by_magnitude():
    from yandex_report.club_aggregate import _fmt_val, _rank_param
    # знак в отображении: «+» для завышения, «−» для занижения
    assert _fmt_val(0.05, "%", True, signed=True) == "+5%"
    assert _fmt_val(-0.02, "%", True, signed=True) == "-2%"
    # место — по модулю: −2% ближе к нулю, чем +5% → 1-е место
    clubs = [("A", {"deviation": 0.05}), ("B", {"deviation": -0.02})]
    _stats, rows = _rank_param(clubs, "deviation", "asc")
    by_place = {name: place for place, name, _v, _g in rows}
    assert by_place["B"] == 1 and by_place["A"] == 2


def test_byn_rate_applies_only_to_minsk():
    from yandex_report.club_aggregate import parse_aggregate_clubs
    files = [("Динамо Минск.xlsx", _make_raw()), ("Сибирь.xlsx", _make_raw())]
    base = parse_aggregate_clubs(aggregate_clubs(files, capacity=200))
    conv = parse_aggregate_clubs(aggregate_clubs(files, capacity=200, byn_to_rub=27.0))
    # деньги Минска переведены по курсу, у Сибири — без изменений
    assert conv["Динамо Минск"]["income_total"] == pytest.approx(
        base["Динамо Минск"]["income_total"] * 27.0)
    assert conv["Сибирь"]["income_total"] == base["Сибирь"]["income_total"]


def test_read_calc_metrics_from_raschety_sheet():
    from yandex_report.club_aggregate import read_calc_metrics
    wb = openpyxl.load_workbook(BytesIO(_make_raw()))
    ws = wb.create_sheet("Расчеты")
    for i, (lbl, val) in enumerate([
        ("Средняя цена платного билета (регулярка)", 1000),
        ("Доля бесплатных билетов", 0.1),
        ("Доля онлайн-продаж", 0.9),
        ("Доход — всего за сезон", 5_000_000),
        ("Агентская комиссия (средняя по агентам)", 0.05),
        ("Расхождение факта с заявленным, %", -0.02),
    ], start=1):
        ws.cell(i, 2, lbl); ws.cell(i, 3, val)
    buf = BytesIO(); wb.save(buf)

    m = read_calc_metrics(buf.getvalue(), capacity=200)
    assert m["price_reg"] == 1000 and m["free_share"] == 0.1 and m["online_share"] == 0.9
    assert m["income_total"] == 5_000_000 and m["commission"] == 0.05
    assert m["deviation"] == -0.02
    # без листа «Расчеты» — None (агрегат посчитает сам)
    assert read_calc_metrics(_make_raw()) is None


def test_too_many_files_raises():
    files = [(f"c{i}.xlsx", _make_raw()) for i in range(MAX_CLUBS + 1)]
    with pytest.raises(TicketError):
        aggregate_clubs(files)


def test_empty_raises():
    with pytest.raises(TicketError):
        aggregate_clubs([])
