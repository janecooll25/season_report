from __future__ import annotations

from io import BytesIO

import openpyxl
from docx import Document

from test_ticket_metrics import _make_raw
from yandex_report.club_report import (
    build_club_report, list_clubs, parse_region_income, region_rank,
)

# сезоны в блоке параметра: 18/19 … 25/26 (индексы 0..7)
SEASONS = ["18/19", "19/20", "20/21", "21/22", "22/23", "23/24", "24/25", "25/26"]
# (стартовый столбец, значения по клубам для 23/24, 24/25)
PARAMS = {
    2: {"Альфа": (2, 2), "Бета": (3, 3), "Гамма": (1, 1)},        # отклонение (градация)
    10: {"Альфа": (0.80, 0.90), "Бета": (0.50, 0.60), "Гамма": (0.30, 0.40)},  # онлайн
    18: {"Альфа": (800, 900), "Бета": (600, 650), "Гамма": (400, 420)},        # цена
    26: {"Альфа": (0.20, 0.15), "Бета": (0.30, 0.25), "Гамма": (0.40, 0.35)},  # безвозм.
}


def _make_monitoring() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Чек-лист мониторинг"
    ws["A5"] = "Отклонение данных официальных протоколов матчей от данных выгрузок"
    ws["J5"] = "Доля розничных продаж билетов через интернет"
    ws["R5"] = "Средняя цена коммерческой реализации билетов в рамках Регулярного чемпионата"
    ws["Z5"] = "Процент билетов и абонементов, распространяемых на безвозмездной основе"
    for start in (2, 10, 18, 26):
        for i, s in enumerate(SEASONS):
            ws.cell(6, start + i, s)
    clubs = ["Альфа", "Бета", "Гамма"]
    for ri, club in enumerate(clubs, start=7):
        ws.cell(ri, 1, club)
        for start, vals in PARAMS.items():
            v2324, v2425 = vals[club]
            ws.cell(ri, start + 5, v2324)  # 23/24
            ws.cell(ri, start + 6, v2425)  # 24/25
    ws.cell(10, 1, "Характеристика")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_grade_thresholds_match_monitoring_colors():
    from yandex_report.club_report import PARAM_FREE, PARAM_ONLINE, _grade_by_value
    # онлайн: ≥85% хорошо, 78–85% удовл., ниже неуд.
    assert _grade_by_value(PARAM_ONLINE, 0.92) == "Хорошие показатели"
    assert _grade_by_value(PARAM_ONLINE, 0.80) == "Удовлетворительные показатели"
    assert _grade_by_value(PARAM_ONLINE, 0.761) == "Неудовлетворительные показатели"
    # бесплатные: ≤11% хорошо, 11–20% удовл., выше неуд.
    assert _grade_by_value(PARAM_FREE, 0.08) == "Хорошие показатели"
    assert _grade_by_value(PARAM_FREE, 0.177) == "Удовлетворительные показатели"
    assert _grade_by_value(PARAM_FREE, 0.36) == "Неудовлетворительные показатели"


def test_deviation_oversold_recommends_reselling_seat():
    from yandex_report.club_report import PARAM_DEV, _describe
    # продали больше, чем пришло (d<0): совет про реализацию места абонемента
    char, rec, _g = _describe(PARAM_DEV, "grade", "asc", -0.05, None, {}, "X")
    assert "больше билетов" in char.lower()
    assert "реализации места владельца абонемента" in rec.lower()
    # по протоколу пришло больше (d>0): без этого совета (юр. абзац)
    char2, rec2, _g2 = _describe(PARAM_DEV, "grade", "asc", 0.05, None, {}, "X")
    assert "реализации места владельца абонемента" not in rec2.lower()


def test_deviation_graded_by_sign():
    from yandex_report.club_report import PARAM_DEV, _describe
    g = lambda cur: _describe(PARAM_DEV, "grade", "asc", cur, None, {}, "X")[2]
    # ≈0 — точное совпадение маловероятно → удовлетворительно (предупреждение)
    assert g(0.0) == "Удовлетворительные показатели"
    # протокол превышает билеты (d>0) — плохо, даже на 1–2%
    assert g(0.01) == "Неудовлетворительные показатели"
    assert g(0.02) == "Неудовлетворительные показатели"
    # продали больше, чем пришло (d<0) — приемлемо (пустые места)
    assert g(-0.01) == "Хорошие показатели"
    assert g(-0.10) == "Хорошие показатели"
    # готовая градация 1/2/3 из мониторинга — как есть
    assert g(3) == "Хорошие показатели"
    assert g(1) == "Неудовлетворительные показатели"


def test_price_high_occupancy_from_90pct():
    from yandex_report.club_report import PARAM_PRICE, _describe
    lg = {"A": 900, "B": 700}
    # 68% — не «стабильно высокая»: удовлетворительно + совет снизить цену
    _c, rec, grade = _describe(PARAM_PRICE, "money", "desc", 800, 750, lg, "X", fill=0.68)
    assert grade == "Удовлетворительные показатели"
    assert "стабильно высок" not in rec.lower() and "68%" in rec
    # 93% — стабильно высокая: хорошо
    _c2, rec2, grade2 = _describe(PARAM_PRICE, "money", "desc", 800, 750, lg, "X", fill=0.93)
    assert grade2 == "Хорошие показатели" and "стабильно высок" in rec2.lower()
    # заполняемость неизвестна — не заявляем высокую
    _c3, rec3, grade3 = _describe(PARAM_PRICE, "money", "desc", 800, 750, lg, "X", fill=None)
    assert "стабильно высок" not in rec3.lower()


def test_commission_mode_wordings():
    from yandex_report.club_report import PARAM_COMMISSION, _describe
    # 1) не реализует через агентов → благоприятно (зелёный)
    char, _rec, grade = _describe(PARAM_COMMISSION, "commission", "asc", None, None, {},
                                  "X", commission_mode="none")
    assert "не реализует билеты через агентов" in char.lower()
    assert grade == "Хорошие показатели"
    # 2) не раскрывает комиссию → без градации (нельзя оценить)
    char2, _r2, grade2 = _describe(PARAM_COMMISSION, "commission", "asc", None, None, {},
                                   "X", commission_mode="undisclosed")
    assert "не раскрывает размер" in char2.lower()
    assert grade2 == ""


def test_shanghai_uses_kunlun_history_row():
    import openpyxl
    from yandex_report.club_report import _club_rows, _monitor_row, _norm_club
    # Шанхай = переименованный Куньлунь → одно нормализованное имя
    assert _norm_club("Шанхайские Драконы") == _norm_club("Куньлунь Ред Стар")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Чек-лист мониторинг"
    ws["A5"] = "Доля розничных продаж билетов через интернет"
    for i, s in enumerate(SEASONS):
        ws.cell(6, 10 + i, s)
    ws["A7"] = "Куньлунь Ред Стар"
    ws.cell(7, 10 + 6, 0.7)          # у Куньлуня есть история
    ws["A8"] = "Шанхайские Драконы"  # новая строка без данных
    ws["A9"] = "Характеристика"
    rows = _club_rows(ws)
    # для Шанхая берём строку Куньлуня (там данные), а не пустую свою
    assert _monitor_row(ws, rows, "Шанхайские Драконы") == 7


def test_agg_lookup_matches_dynamo_name_variants():
    from yandex_report.club_report import _agg_lookup
    agg = {"Динамо М": {"online_share": 0.99}, "Динамо Мн": {"online_share": 0.77}}
    # чек-лист называет клубы полнее — сопоставляем по нормализованному имени
    assert _agg_lookup(agg, "Динамо Москва")["online_share"] == 0.99
    assert _agg_lookup(agg, "Динамо Минск")["online_share"] == 0.77
    assert _agg_lookup(agg, "Сибирь") == {}


def test_fill_monitoring_writes_season_values_and_colors():
    from yandex_report.club_aggregate import aggregate_clubs
    from yandex_report.club_report import fill_monitoring
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Чек-лист мониторинг"
    ws["B5"] = "Отклонение данных официальных протоколов"
    ws["J5"] = "Доля розничных продаж билетов через интернет"
    ws["R5"] = "Средняя цена коммерческой реализации билетов"
    ws["Z5"] = "Процент билетов, распространяемых на безвозмездной основе"
    for start in (2, 10, 18, 26):
        for i, s in enumerate(SEASONS):
            ws.cell(6, start + i, s)
    ws.cell(7, 1, "Сибирь")
    ws.cell(8, 1, "Характеристика")
    buf = BytesIO(); wb.save(buf)

    agg = aggregate_clubs([("Сибирь.xlsx", _make_raw())], capacity=200)
    out = fill_monitoring(buf.getvalue(), agg)
    ws2 = openpyxl.load_workbook(BytesIO(out))["Чек-лист мониторинг"]
    # онлайн 25/26 (col 10+7=17) = 300/400 = 0.75, ячейка залита
    assert round(ws2.cell(7, 17).value, 2) == 0.75
    assert ws2.cell(7, 17).fill.patternType == "solid"
    # цена 25/26 (col 25) заполнена числом
    assert isinstance(ws2.cell(7, 25).value, (int, float))


def test_list_clubs():
    assert list_clubs(_make_monitoring()) == ["Альфа", "Бета", "Гамма"]


def _rows_by_param(out):
    doc = Document(BytesIO(out))
    return {r.cells[1].text: (r.cells[2].text, r.cells[3].text)
            for r in doc.tables[0].rows[1:]}


def test_report_from_monitoring_has_place_and_dynamics():
    out = build_club_report(_make_monitoring(), "Альфа", season="24/25", prev_season="23/24")
    doc = Document(BytesIO(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "«Альфа»" in all_text
    assert "2024/2025" in all_text
    rows = _rows_by_param(out)
    # цена: Альфа 900 — лучшая из трёх → 1 место, среднее (900+650+420)/3≈657
    price_char = next(v[0] for k, v in rows.items() if "Средняя цена" in k)
    assert "900 руб" in price_char and "1 месте" in price_char
    # онлайн: динамика 80% → 90%
    online_rec = next(v[1] for k, v in rows.items() if "интернет" in k)
    assert "с 80% до 90%" in online_rec


def test_deviation_grade_from_monitoring_not_percent():
    out = build_club_report(_make_monitoring(), "Бета", season="24/25", prev_season="23/24")
    rows = _rows_by_param(out)
    dev_char = next(v[0] for k, v in rows.items() if "Отклонение" in k)
    # градация 3 → «хорошие показатели», без «300%»
    assert "300%" not in dev_char
    assert "хорошие показатели" in dev_char.lower()


def _make_income() -> bytes:
    """Мини-файл СДД: годовой столбец 2025 в колонке 3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "СДД_субъекты"
    ws.cell(6, 2, "2025 год 3)")
    ws.cell(7, 2, "I кв.")
    ws.cell(7, 3, "год")
    data = [
        ("Российская Федерация", 74932),          # исключается
        ("Сибирский федеральный округ 2)", 58485),  # исключается
        ("Новосибирская область", 68111),         # Сибирь
        ("Омская область", 54959),                # Авангард
        ("г. Москва", 165866),                    # ЦСКА/Спартак/Динамо М
        ("Республика Башкортостан", 54198),       # Салават Юлаев
    ]
    for i, (name, val) in enumerate(data, start=8):
        ws.cell(i, 1, name)
        ws.cell(i, 3, val)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_region_income_and_rank():
    inc = parse_region_income(_make_income(), year="2025")
    assert "Новосибирская область" in inc and inc["Новосибирская область"] == 68111
    assert "Российская Федерация" not in inc          # РФ исключена
    assert not any("федеральн" in k.lower() for k in inc)  # ФО исключены
    # среди 4 регионов присутствия (Новосиб, Омск, Москва, Башкортостан)
    region, place, n, val = region_rank(inc, "Сибирь")
    assert region == "Новосибирская область" and n == 4
    assert place == 2 and val == 68111  # Москва 1-я, Новосибирск 2-й


def _make_monitoring_real() -> bytes:
    """Мониторинг с реальными названиями клубов (для привязки к регионам)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Чек-лист мониторинг"
    ws["R5"] = "Средняя цена коммерческой реализации билетов в рамках Регулярного чемпионата"
    for i, s in enumerate(SEASONS):
        ws.cell(6, 18 + i, s)
    for ri, (club, v2324, v2425) in enumerate(
            [("Авангард", 880, 960), ("Сибирь", 906, 1019), ("ЦСКА", 1500, 1600)], start=7):
        ws.cell(ri, 1, club)
        ws.cell(ri, 18 + 5, v2324)
        ws.cell(ri, 18 + 6, v2425)
    ws.cell(10, 1, "Характеристика")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_region_note_in_price_recommendation():
    out = build_club_report(_make_monitoring_real(), "Авангард", season="24/25",
                            prev_season="23/24", region_income_bytes=_make_income())
    rows = _rows_by_param(out)
    price_rec = next(v[1] for k, v in rows.items() if "Средняя цена" in k)
    assert "Омская область" in price_rec and "среднедушев" in price_rec.lower()


def test_commission_row_and_aggregate_source():
    from yandex_report.club_aggregate import aggregate_clubs
    agg = aggregate_clubs([("Авангард.xlsx", _make_raw()),
                           ("Сибирь.xlsx", _make_raw())], capacity=200)
    out = build_club_report(_make_monitoring_real(), "Авангард", season="25/26",
                            prev_season="24/25", aggregate_bytes=agg)
    rows = _rows_by_param(out)
    # строка агентской комиссии присутствует
    assert any("комисси" in k.lower() for k in rows)
    # значение онлайна взято из агрегата (_make_raw: 300 из 400 = 75%)
    online = next(v[0] for k, v in rows.items() if "интернет" in k)
    assert "75%" in online


def test_current_season_from_ticket_file():
    # 25/26 в мониторинге пуст → значения берутся из билетного файла
    out = build_club_report(_make_monitoring(), "Альфа", season="25/26",
                            prev_season="24/25", ticket_bytes=_make_raw(), capacity=200)
    rows = _rows_by_param(out)
    online_char = next(v[0] for k, v in rows.items() if "интернет" in k)
    # _make_raw: онлайн 300 из 400 = 75%
    assert "75%" in online_char
