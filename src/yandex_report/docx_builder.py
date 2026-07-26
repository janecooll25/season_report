"""Сборка нового docx-отчёта из структуры, данных и сгенерированного текста."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from .report_data import DataBlock
from .structure import COVER_LETTER, SECTIONS, Section, appendix_title


def _add_table(doc: Document, block: DataBlock) -> None:
    if not block.rows:
        doc.add_paragraph("Нет данных за выбранный период.")
        return
    table = doc.add_table(rows=1, cols=len(block.columns))
    table.style = "Light Grid Accent 1"
    for i, col in enumerate(block.columns):
        cell = table.rows[0].cells[i]
        cell.text = col
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for row in block.rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)


def _add_chart(doc: Document, block: DataBlock, chart_type: str, title: str) -> None:
    try:
        from .charts import render_chart

        png = render_chart(block, chart_type, title)
    except Exception:  # noqa: BLE001 — диаграмма не должна ронять отчёт
        png = None
    if png:
        doc.add_picture(BytesIO(png), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _assemble(
    *,
    season: str,
    counter_id: str,
    generated_at: str,
    blocks: dict[str, DataBlock],
    prose: dict[str, str],
    charts: dict[str, str] | None = None,
) -> Document:
    charts = charts or {}
    doc = Document()

    # Титул приложения
    title = doc.add_heading(appendix_title(season), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Счётчик Яндекс.Метрики: {counter_id}    ")
    meta.add_run(f"Сформирован: {generated_at}")

    # Сопроводительный абзац
    doc.add_paragraph(COVER_LETTER.format(season=season))

    # Разделы (пропускаем те, для которых нет ни данных, ни текста —
    # актуально для CSV-режима с частичным набором выгрузок)
    for section in SECTIONS:
        block = blocks.get(section.id)
        text = prose.get(section.id, "").strip()
        if not block and not text:
            continue

        doc.add_heading(section.heading.format(season=season), level=1)

        if text:
            for para in text.split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())

        if block:
            if section.table_title:
                cap = doc.add_paragraph()
                cap.add_run(section.table_title).italic = True
            _add_table(doc, block)

            chart_type = (charts.get(section.id) or "").strip()
            if chart_type:
                _add_chart(doc, block, chart_type,
                           section.table_title or section.heading.format(season=season))

    note = doc.add_paragraph()
    note.add_run(
        "Данные приведены из системы Яндекс.Метрика. Разделы по мобильному "
        "приложению и Клубам КХЛ требуют отдельных источников (AppMetrica / "
        "Google Analytics / счётчики Клубов) и в этот отчёт не включены."
    ).italic = True

    return doc


def build_report(
    *,
    season: str,
    counter_id: str,
    generated_at: str,
    blocks: dict[str, DataBlock],
    prose: dict[str, str],
    output_path: str | Path,
    charts: dict[str, str] | None = None,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = _assemble(
        season=season,
        counter_id=counter_id,
        generated_at=generated_at,
        blocks=blocks,
        prose=prose,
        charts=charts,
    )
    doc.save(str(output_path))
    return output_path


def build_report_bytes(
    *,
    season: str,
    counter_id: str,
    generated_at: str,
    blocks: dict[str, DataBlock],
    prose: dict[str, str],
    charts: dict[str, str] | None = None,
) -> bytes:
    """Собирает отчёт в память и возвращает байты .docx (для веб-выдачи)."""
    doc = _assemble(
        season=season,
        counter_id=counter_id,
        generated_at=generated_at,
        blocks=blocks,
        prose=prose,
        charts=charts,
    )
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
